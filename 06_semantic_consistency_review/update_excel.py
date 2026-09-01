#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
将 sql_nl_test_samples_500_0807_result_data.json 的内容写入评测集人工评测表_500条.xlsx
- JSON 视为原始数据源，只写前4列
- 从 _updated.xlsx 的 Col7(修改前内容) 获取旧值用于对比
- NL+SQL无变化 → 整行黄底，有变化 → 白底
- 清空评测列中脚本自动填充的值
"""

import json, re, sys
import openpyxl
from openpyxl.styles import PatternFill

sys.stdout.reconfigure(encoding='utf-8')

JSON_FILE = r'f:\Study\科研任务\论文\评测集\Code\评测集修改\500\0807\sql_nl_test_samples_500_0807_result_data.json'
XLSX_FILE = r'f:\Study\科研任务\论文\评测集\Code\评测集修改\500\0807\excel\评测集人工评测表_500条.xlsx'
BACKUP_FILE = r'f:\Study\科研任务\论文\评测集\Code\评测集修改\500\0807\excel\评测集人工评测表_500条_updated.xlsx'

YELLOW_FILL = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
NO_FILL = PatternFill(fill_type=None)


def extract_old_nl_sql(before_text):
    """从 Col7 中提取旧 NL 和 SQL"""
    old_nl = ''
    old_sql = ''
    if not before_text:
        return old_nl, old_sql
    for line in before_text.split('\n'):
        if line.startswith('[NL] '):
            old_nl = line[5:].strip()
        elif line.startswith('[SQL] '):
            old_sql = line[5:].strip()
    return old_nl, old_sql


def main():
    # 1. 读取 JSON
    print(f"读取 JSON: {JSON_FILE}")
    with open(JSON_FILE, 'r', encoding='utf-8') as f:
        json_data = json.load(f)
    json_index = {item['sample_id']: item for item in json_data}
    print(f"  {len(json_data)} 条")

    # 2. 读取备份 Excel（获取旧值）
    old_values = {}
    try:
        wb_old = openpyxl.load_workbook(BACKUP_FILE)
        ws_old = wb_old.active
        for row in range(2, ws_old.max_row + 1):
            sid = ws_old.cell(row=row, column=2).value
            if sid is None:
                continue
            try:
                sid = int(sid)
            except (ValueError, TypeError):
                pass
            before_text = ws_old.cell(row=row, column=7).value or ''
            nl, sql = extract_old_nl_sql(before_text)
            if nl or sql:
                old_values[sid] = (nl, sql)
        wb_old.close()
        print(f"  从备份恢复 {len(old_values)} 条旧值")
    except Exception as e:
        print(f"  [WARN] 无法读取备份: {e}")

    # 3. 读取当前 Excel
    print(f"读取 Excel: {XLSX_FILE}")
    wb = openpyxl.load_workbook(XLSX_FILE)
    ws = wb.active
    print(f"  {ws.max_row - 1} 条, {ws.max_column} 列")

    # 4. 列映射
    headers = {}
    for col in range(1, ws.max_column + 1):
        headers[ws.cell(row=1, column=col).value] = col

    col_sid   = headers.get('sample_id')
    col_nl    = headers.get('自然语言Query (natural_language)')
    col_sql   = headers.get('原始SQL语句 (sql)')
    col_seq   = headers.get('评测序号')
    col_mod   = headers.get('是否修改')
    col_mb    = headers.get('修改前内容')
    col_ma    = headers.get('修改后内容')

    # 5. 逐行处理
    unchanged = 0
    changed = 0

    for row in range(2, ws.max_row + 1):
        sid = ws.cell(row=row, column=col_sid).value
        if sid is None:
            continue
        try:
            sid = int(sid)
        except (ValueError, TypeError):
            pass

        json_item = json_index.get(sid)
        if json_item is None:
            continue

        new_nl  = (json_item.get('natural_language', '')).strip()
        new_sql = (json_item.get('sql', '')).strip()

        # 判断旧值：优先从备份的Col7提取，否则用当前Excel值
        if sid in old_values:
            old_nl, old_sql = old_values[sid]
        else:
            old_nl = (ws.cell(row=row, column=col_nl).value or '').strip()
            old_sql = (ws.cell(row=row, column=col_sql).value or '').strip()

        is_unchanged = (old_nl == new_nl) and (old_sql == new_sql)

        # 写前4列
        if col_seq:
            ws.cell(row=row, column=col_seq, value=row - 1)
        ws.cell(row=row, column=col_sid, value=sid)
        ws.cell(row=row, column=col_nl, value=new_nl)
        ws.cell(row=row, column=col_sql, value=new_sql)

        # 清空评测列中被脚本自动填充的列
        if col_mod:
            ws.cell(row=row, column=col_mod, value=None)
        if col_mb:
            ws.cell(row=row, column=col_mb, value=None)
        if col_ma:
            ws.cell(row=row, column=col_ma, value=None)

        # 涂色
        fill = YELLOW_FILL if is_unchanged else NO_FILL
        for c in range(1, ws.max_column + 1):
            ws.cell(row=row, column=c).fill = fill

        if is_unchanged:
            unchanged += 1
        else:
            changed += 1

    # 6. 保存
    print(f"\n保存: {XLSX_FILE}")
    try:
        wb.save(XLSX_FILE)
        print("完成!")
    except PermissionError:
        alt = XLSX_FILE.replace('.xlsx', '_NEW.xlsx')
        print(f"原文件被占用，保存到: {alt}")
        wb.save(alt)

    print("\n" + "=" * 60)
    print(f"  总条目:      {unchanged + changed}")
    print(f"  无变化(黄底): {unchanged} 条")
    print(f"  有变化(白底): {changed} 条")
    print("=" * 60)


if __name__ == '__main__':
    main()
